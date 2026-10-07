"""Gera entregas/projeto-escrito.docx no formato do modelo UniGestor (docs/fonte/relatorio_extensao_projeto word.docx)
e os anexos. O conteúdo fica em scripts/conteudo_projeto.py; edite lá e rode: python3 scripts/gerar_relatorio.py"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
import conteudo_projeto as C  # noqa: E402

RAIZ = os.path.join(os.path.dirname(__file__), "..")
CINZA = RGBColor(0x44, 0x44, 0x44)
CINZA_CLARO = RGBColor(0x66, 0x66, 0x66)
ASSETS = os.path.join(RAIZ, "assets")


def cabecalho_logos(doc, altura=1.4):
    """Logo da instituição à esquerda e selo do projeto à direita, como na capa do pré-projeto."""
    t = doc.add_table(rows=1, cols=2)
    esq, dir_ = t.rows[0].cells
    esq.paragraphs[0].add_run().add_picture(os.path.join(ASSETS, "logo-unicesumar.png"), height=Cm(altura * 0.75))
    dir_.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    dir_.paragraphs[0].add_run().add_picture(os.path.join(ASSETS, "selo-efeito-rebote.jpg"), height=Cm(altura * 1.4))
    esq.vertical_alignment = 1  # centro
    larguras(t, (9.0, 9.0))


def base(doc):
    cp = doc.core_properties
    cp.author = cp.last_modified_by = "Acadêmicos do 3º semestre noturno – Turma B"
    cp.comments = ""
    cp.title = "Projeto de Extensão Efeito Rebote"
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5), Cm(1.4), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(9.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(0)


def run(p, texto, bold=False, size=None, color=None, italic=False):
    r = p.add_run(texto)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return r


def borda_inferior(p, cor="BBBBBB"):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", cor)):
        e.set(qn(k), v)
    b.append(e)
    pPr.append(b)


def bordas_tabela(t, cor="DDDDDD"):
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", cor)):
            e.set(qn(k), v)
        b.append(e)
    tblPr.append(b)


def larguras(t, cms):
    """Fixa a grade da tabela (Word e LibreOffice ignoram só a largura da célula)."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, cm in zip(grid.findall(qn("w:gridCol")), cms):
        gc.set(qn("w:w"), str(int(cm * 567)))
    for row in t.rows:
        for cel, cm in zip(row.cells, cms):
            cel.width = Cm(cm)


def secao(doc, titulo):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run(p, titulo, bold=True, size=10.5)
    borda_inferior(p)


def campos(doc, pares, largura_rotulo=4.2):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    bordas_tabela(t, "EEEEEE")
    for rotulo, valor in pares:
        c1, c2 = t.add_row().cells
        c1.width, c2.width = Cm(largura_rotulo), Cm(18 - largura_rotulo)
        run(c1.paragraphs[0], rotulo, bold=True, color=CINZA)
        paragrafos = valor if isinstance(valor, list) else [valor]
        for i, texto in enumerate(paragrafos):
            p = c2.paragraphs[0] if i == 0 else c2.add_paragraph()
            if isinstance(texto, dict):  # figura: {"img", "legenda", "largura"}
                p.paragraph_format.space_before = Pt(8)
                run(p, texto["legenda"], bold=True, size=8.5, color=CINZA)
                q = c2.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                q.add_run().add_picture(os.path.join(ASSETS, texto["img"]), width=Cm(texto.get("largura", 12)))
                f = c2.add_paragraph()
                run(f, "Fonte: Elaborado pelos autores (2026).", size=8, color=CINZA_CLARO)
                continue
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if i:
                p.paragraph_format.space_before = Pt(4)
            run(p, texto, color=CINZA)
    larguras(t, (largura_rotulo, 18 - largura_rotulo))
    return t


def cronograma(doc):
    t = doc.add_table(rows=1, cols=3)
    bordas_tabela(t)
    larg = (Cm(1.8), Cm(14.2), Cm(2.0))
    for cel, txt, w in zip(t.rows[0].cells, ("Encontro", "Descrição", "Carga (h)"), larg):
        cel.width = w
        run(cel.paragraphs[0], txt, bold=True)
    total = 0
    for n, enc in enumerate(C.CRONOGRAMA, 1):
        c = t.add_row().cells
        for cel, w in zip(c, larg):
            cel.width = w
        run(c[0].paragraphs[0], str(n), color=CINZA)
        c[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = c[1].paragraphs[0]
        run(p, f"ENCONTRO {n} – {enc['titulo']}", bold=True, color=CINZA)
        for rotulo, chave in (("Etapa da metodologia", "etapa"), ("Objetivos", "objetivos"),
                              ("Atividades", "atividades"), ("Entrega do estudante", "entrega")):
            q = c[1].add_paragraph()
            q.paragraph_format.space_before = Pt(2)
            run(q, f"{rotulo}: ", bold=True, color=CINZA)
            run(q, enc[chave], color=CINZA)
        c[1].add_paragraph()
        run(c[2].paragraphs[0], f"{enc['carga']:.2f}" if enc["carga"] else "", color=CINZA)
        c[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        total += enc["carga"] or 0
    larguras(t, (1.8, 14.2, 2.0))
    return total


def gerar_relatorio():
    doc = Document()
    base(doc)
    cabecalho_logos(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run(p, "Relatório da Atividade Extensionista", bold=True, size=16)
    p = doc.add_paragraph()
    run(p, C.CABECALHO, size=8, color=CINZA_CLARO)

    secao(doc, "1) Identificação")
    campos(doc, C.IDENTIFICACAO)

    total = sum(e["carga"] or 0 for e in C.CRONOGRAMA)
    secao(doc, "2) Períodos e vagas")
    pares = [(k, (f"{total:g} h" if v == "{CARGA_TOTAL}" else v)) for k, v in C.PERIODOS]
    campos(doc, pares + [("Comunidade participante", C.COMUNIDADE)])

    secao(doc, "3) Dimensão pedagógica")
    campos(doc, C.DIMENSAO_PEDAGOGICA)

    secao(doc, "4) Cronograma")
    cronograma(doc)

    secao(doc, "5) Anexos")
    t = doc.add_table(rows=1, cols=2)
    bordas_tabela(t)
    run(t.rows[0].cells[0].paragraphs[0], "Arquivo", bold=True)
    run(t.rows[0].cells[1].paragraphs[0], "Código", bold=True)
    for nome, cod in C.ANEXOS:
        c = t.add_row().cells
        run(c[0].paragraphs[0], nome, color=CINZA)
        run(c[1].paragraphs[0], cod, color=CINZA)
    larguras(t, (15.0, 3.0))

    destino = os.path.join(RAIZ, "entregas", "projeto-escrito.docx")
    doc.save(destino)
    return destino, total


def gerar_anexo(nome_arquivo, titulo, blocos):
    """blocos: lista de (tipo, texto) com tipo em {'h', 'p', 'li'}."""
    doc = Document()
    base(doc)
    doc.styles["Normal"].font.size = Pt(11)
    cabecalho_logos(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, titulo, bold=True, size=13)
    p.paragraph_format.space_after = Pt(10)
    for tipo, texto in blocos:
        if tipo == "h":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            run(p, texto, bold=True)
        elif tipo == "li":
            p = doc.add_paragraph(style="List Bullet")
            run(p, texto)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            run(p, texto)
    destino = os.path.join(RAIZ, "entregas", nome_arquivo)
    doc.save(destino)
    return destino


if __name__ == "__main__":
    d, total = gerar_relatorio()
    print(f"ok {d} (carga total {total:g} h)")
    print("ok", gerar_anexo("anexo-1-regulamento-acao-arrecadacao.docx", C.REGULAMENTO_TITULO, C.REGULAMENTO))
    print("ok", gerar_anexo("anexo-2-texto-apresentacao.docx", C.APRESENTACAO_TITULO, C.APRESENTACAO))
