"""Gera a organização das equipes da turma (funções, vagas e líderes) em entregas/equipes/:
organizacao-equipes.md, organizacao-equipes.html (-> PDF com scripts/renderizar.mjs) e planilha-equipes.xlsx.
Os tamanhos são parametrizados: ajuste ALUNOS e as vagas de cada função e rode de novo.
Uso: python3 scripts/gerar_equipes.py"""
import html
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.join(os.path.dirname(__file__), "..")
SAIDA = os.path.join(RAIZ, "entregas", "equipes")
ALUNOS = 80  # "cerca de 80 alunos, dos períodos matutino e noturno" (docs/projeto.md)

# Equipe transversal: todos participam, além da equipe fixa de cada um
TRANSVERSAL = {
    "nome": "Arrecadação e Captação",
    "membros": "TODOS os acadêmicos",
    "missao": "Garantir que cada acadêmico cumpra sua parte na ação de arrecadação e ampliar a coleta de itens com parceiros da comunidade.",
    "lider": [
        "coordenar a entrega dos blocos de bilhetes e registrar qual numeração ficou com cada acadêmico;",
        "reunir-se às segundas com os líderes das demais equipes, que atuam como pontos focais e acompanham os próprios membros;",
        "manter a lista de parceiros (comércios, igrejas, delegacias) e das caixas de coleta, sempre com autorização do responsável pelo local e aprovação prévia da prof.ª Camila;",
        "avisar a equipe de Triagem sobre cada caixa nova a recolher.",
    ],
    "todos": [
        "vender os 30 bilhetes do seu bloco e repassar R$ 150,00 por PIX, com os 30 canhotos, até 30/10, às 23h59;",
        "doar e trazer itens da lista aceita (escova, creme dental de até 100 g, aparelho de barbear descartável, detergente transparente e sabão em pó);",
        "indicar possíveis parceiros ao líder (não fechar parceria por conta própria);",
        "compartilhar os posts oficiais, sem pedir doações antes da liberação.",
    ],
}

# (nome, missão, [(função, vagas, tarefas)]) — a 1ª função é sempre a liderança
EQUIPES = [
    ("Comunicação e Redes Sociais",
     "Manter o Instagram @efeitorebote.oficial ativo (3 posts por semana) e explicar o projeto à sociedade.",
     [("Líder", 1, "fecha o calendário editorial, envia cada peça à prof.ª Camila e só publica o que foi aprovado; ponto focal dos membros na arrecadação."),
      ("Vice-líder (outro período)", 1, "substitui o líder e repassa as decisões ao outro período."),
      ("Roteiro e legendas", 2, "textos dos posts, legendas e roteiros de vídeo e áudio, em linguagem acessível."),
      ("Design", 3, "artes, carrosséis e stories na identidade visual do projeto."),
      ("Audiovisual", 2, "gravação e edição de vídeos e áudios; transmissão ao vivo do sorteio, junto com Eventos."),
      ("Engajamento", 2, "stories diários, resposta a comentários e mensagens, métricas semanais para o relatório."),
      ("Registro de imagens", 2, "fotos e vídeos dos eventos, somente com autorização e nunca de pessoas privadas de liberdade."),
      ("Arquivo de aprovações", 1, "guarda cada versão enviada à professora e a resposta recebida.")]),
    ("Criatividade e Organização de Eventos",
     "Planejar e executar os eventos presenciais do projeto.",
     [("Líder", 1, "cronograma dos eventos, pedidos de autorização (campus e unidades) e divisão das escalas; ponto focal dos membros na arrecadação."),
      ("Vice-líder (outro período)", 1, "substitui o líder e organiza as escalas do outro período."),
      ("Cenário temático no pátio", 6, "montagem, exposição dos itens aceitos, caixa de coleta e escala de atendimento ao público."),
      ("Café da manhã com as famílias", 6, "doações de alimentos, regras da unidade (alimentos, embalagens, revista), acolhimento das crianças; sem registro de imagens."),
      ("Sorteio", 3, "urna, conferência dos canhotos com o Financeiro, ata e testemunhas; apoio à transmissão ao vivo."),
      ("Visitas técnicas (04 e 05/11)", 5, "transporte, lista de presença, grupos de até 50 alunos, checklist das normas de segurança e conduta.")]),
    ("Relatório Final e Documentação",
     "Registrar tudo o que o projeto fizer e produzir o relatório final no modelo da professora.",
     [("Líder", 1, "estrutura do relatório, prazos internos e envio à professora; ponto focal dos membros na arrecadação."),
      ("Vice-líder (outro período)", 1, "substitui o líder e coleta os registros do outro período."),
      ("Atas e frequência", 2, "ata de cada reunião de segunda e lista de presença dos encontros."),
      ("Evidências", 2, "reúne fotos autorizadas, métricas das redes, planilhas e termos de entrega, organizados por ação."),
      ("Pesquisa e referências", 2, "dados e fontes sobre o sistema prisional, com citações conferidas (ABNT)."),
      ("Redação", 3, "escreve as seções do relatório a partir das evidências."),
      ("Revisão e formatação", 1, "revisão final, padronização no modelo e conferência das citações.")]),
    ("Financeiro e Prestação de Contas",
     "Atuar como a comissão financeira prevista no regulamento: controlar cada real que entra e sai.",
     [("Líder", 1, "responde pela planilha de controle e pelo relatório de transparência; ponto focal dos membros na arrecadação."),
      ("Conferência de repasses", 2, "confere extrato × planilha × canhotos de cada bloco (dupla conferência)."),
      ("Compras e notas fiscais", 1, "orçamentos, compra do prêmio e dos itens, sempre com nota fiscal."),
      ("Despesas e comprovantes", 1, "lança as saídas e arquiva os comprovantes."),
      ("Transparência", 1, "prepara o relatório de transparência para a professora e a turma.")]),
    ("Triagem e Aferição dos Itens",
     "Recolher, conferir, contar e preparar os itens para a entrega nas unidades.",
     [("Líder", 1, "escala de recolhimento, local de armazenamento e termo de entrega por unidade; ponto focal dos membros na arrecadação."),
      ("Vice-líder (outro período)", 1, "substitui o líder e coordena a triagem no outro período."),
      ("Recolhimento no campus", 4, "esvazia as caixas de coleta do campus toda semana."),
      ("Recolhimento nos parceiros", 5, "recolhe semanalmente as caixas dos comércios, igrejas e delegacias."),
      ("Triagem e conformidade", 8, "confere se o item está na lista, lacrado e dentro das regras (creme dental de até 100 g, aparelho descartável de duas lâminas); separa o que não serve."),
      ("Contagem e registro", 3, "lança na planilha a quantidade por tipo de item."),
      ("Kits e armazenamento", 4, "embala e etiqueta os itens por unidade (PEM, CCM, CPIM) para a entrega.")]),
]

total = sum(v for _, _, fs in EQUIPES for _, v, _ in fs)
assert total == ALUNOS, f"vagas ({total}) diferente de ALUNOS ({ALUNOS})"
os.makedirs(SAIDA, exist_ok=True)

REGRAS = [
    "Cada acadêmico participa de Arrecadação e Captação e de exatamente UMA das cinco equipes fixas.",
    "Cada equipe fixa deve ter membros dos dois períodos; quando há vice-líder, ele é do período oposto ao do líder.",
    "Os líderes das cinco equipes, o líder de Arrecadação e Captação e a prof.ª Camila formam o comitê de líderes, que se reúne às segundas.",
    "Nada é publicado, comprado ou combinado com terceiros sem aprovação da prof.ª Camila.",
    "Quem é titular da conta da ação de arrecadação não deve ser o único a conferir os repasses: a conferência é sempre feita por duas pessoas.",
]

# ---------- Markdown ----------
md = ["# Efeito Rebote: organização das equipes", "",
      f"Estrutura para cerca de {ALUNOS} acadêmicos (matutino e noturno). Nomes a preencher na `planilha-equipes.xlsx`.", "",
      "## Regras gerais", ""] + [f"- {r}" for r in REGRAS] + ["",
      f"## 0. {TRANSVERSAL['nome']} ({TRANSVERSAL['membros']})", "", TRANSVERSAL["missao"], "",
      "**Líder:** ______________________", "", "Funções do líder:"] + [f"- {t}" for t in TRANSVERSAL["lider"]] + ["",
      "Funções de todos:"] + [f"- {t}" for t in TRANSVERSAL["todos"]] + [""]
for i, (nome, missao, fs) in enumerate(EQUIPES, 1):
    n = sum(v for _, v, _ in fs)
    md += [f"## {i}. {nome} ({n} membros)", "", missao, "", "| Função | Vagas | O que faz | Nome(s) |", "|---|---|---|---|"]
    md += [f"| {f} | {v} | {t} | |" for f, v, t in fs] + [""]
md += ["## Resumo", "", "| Equipe | Membros |", "|---|---|"]
md += [f"| {TRANSVERSAL['nome']} | todos |"] + [f"| {n} | {sum(v for _, v, _ in fs)} |" for n, _, fs in EQUIPES]
md += [f"| **Total nas equipes fixas** | **{total}** |", ""]
open(os.path.join(SAIDA, "organizacao-equipes.md"), "w").write("\n".join(md))

# ---------- HTML (para PDF) ----------
e = html.escape
h = ["<!doctype html><html lang='pt-BR'><meta charset='utf-8'><title>Organização das equipes</title><style>",
     "@page{size:A4;margin:14mm}body{font-family:Arial,sans-serif;font-size:10.5pt;color:#1a1a1a;margin:0}",
     "h1{color:#1B3A8C;font-size:19pt;margin:0 0 4px}h2{color:#fff;background:#1B3A8C;font-size:12pt;padding:5px 8px;margin:16px 0 6px}",
     ".sub{color:#555;margin:0 0 8px}table{width:100%;border-collapse:collapse;page-break-inside:auto}tr{page-break-inside:avoid}",
     "th,td{border:1px solid #c9cfdc;padding:4px 6px;vertical-align:top;text-align:left}th{background:#e8ecf6}",
     "td.v{text-align:center;width:38px}td.n{width:150px}td.f{width:150px;font-weight:bold}ul{margin:4px 0 4px 18px;padding:0}",
     ".lid{border:1px dashed #1B3A8C;padding:6px 8px;margin:6px 0}.reg{background:#f4f6fb;padding:6px 10px;border-left:4px solid #1B3A8C}",
     "</style><body>",
     "<h1>Efeito Rebote: organização das equipes</h1>",
     f"<p class='sub'>Estrutura para cerca de {ALUNOS} acadêmicos dos períodos matutino e noturno.</p>",
     "<div class='reg'><b>Regras gerais</b><ul>" + "".join(f"<li>{e(r)}</li>" for r in REGRAS) + "</ul></div>",
     f"<h2>0. {e(TRANSVERSAL['nome'])}: {e(TRANSVERSAL['membros'])}</h2><p>{e(TRANSVERSAL['missao'])}</p>",
     "<div class='lid'><b>Líder:</b> ____________________________________</div>",
     "<b>Funções do líder</b><ul>" + "".join(f"<li>{e(t)}</li>" for t in TRANSVERSAL["lider"]) + "</ul>",
     "<b>Funções de todos</b><ul>" + "".join(f"<li>{e(t)}</li>" for t in TRANSVERSAL["todos"]) + "</ul>"]
for i, (nome, missao, fs) in enumerate(EQUIPES, 1):
    n = sum(v for _, v, _ in fs)
    h += [f"<h2>{i}. {e(nome)}: {n} membros</h2><p>{e(missao)}</p>",
          "<table><tr><th>Função</th><th>Vagas</th><th>O que faz</th><th>Nome(s)</th></tr>"]
    h += [f"<tr><td class='f'>{e(f)}</td><td class='v'>{v}</td><td>{e(t)}</td><td class='n'></td></tr>" for f, v, t in fs]
    h += ["</table>"]
h += ["<h2>Resumo</h2><table><tr><th>Equipe</th><th>Membros</th></tr>",
      f"<tr><td>{e(TRANSVERSAL['nome'])}</td><td>todos</td></tr>"]
h += [f"<tr><td>{e(n)}</td><td>{sum(v for _, v, _ in fs)}</td></tr>" for n, _, fs in EQUIPES]
h += [f"<tr><th>Total nas equipes fixas</th><th>{total}</th></tr></table></body></html>"]
open(os.path.join(SAIDA, "organizacao-equipes.html"), "w").write("\n".join(h))

# ---------- Planilha de vagas ----------
CAB, BRANCO = PatternFill("solid", fgColor="1B3A8C"), Font(bold=True, color="FFFFFF")
LIDER = PatternFill("solid", fgColor="FFF3CD")
wb = Workbook()
wb.properties.creator = wb.properties.lastModifiedBy = "Acadêmicos do 3º semestre noturno – Turma B"
ws = wb.active
ws.title = "Vagas"
ws.append(["Equipe", "Função", "Vaga", "Nome", "Período", "Contato"])
for c, w in zip(ws[1], (34, 30, 6, 34, 12, 18)):
    c.fill, c.font = CAB, BRANCO
    ws.column_dimensions[c.column_letter].width = w
ws.freeze_panes = "A2"
ws.append([TRANSVERSAL["nome"], "Líder", 1, "", "", ""])
for c in ws[ws.max_row]:
    c.fill = LIDER
for nome, _, fs in EQUIPES:
    for f, v, _ in fs:
        for k in range(1, v + 1):
            ws.append([nome, f, k, "", "", ""])
            if f == "Líder":
                for c in ws[ws.max_row]:
                    c.fill = LIDER
dv = DataValidation(type="list", formula1='"Matutino,Noturno"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"E2:E{ws.max_row}")
fim = ws.max_row

res = wb.create_sheet("Resumo", 0)
res.append(["Equipe", "Vagas", "Preenchidas", "Matutino", "Noturno", "Líder"])
for c, w in zip(res[1], (38, 8, 12, 10, 10, 30)):
    c.fill, c.font = CAB, BRANCO
    res.column_dimensions[c.column_letter].width = w
for nome in [TRANSVERSAL["nome"]] + [n for n, _, _ in EQUIPES]:
    r = res.max_row + 1
    rng = f"Vagas!$A$2:$A${fim}"
    res.append([nome, f'=COUNTIF({rng},A{r})', f'=COUNTIFS({rng},A{r},Vagas!$D$2:$D${fim},"<>")',
                f'=COUNTIFS({rng},A{r},Vagas!$E$2:$E${fim},"Matutino")', f'=COUNTIFS({rng},A{r},Vagas!$E$2:$E${fim},"Noturno")',
                f'=IFERROR(INDEX(Vagas!$D$2:$D${fim},MATCH(1,INDEX(({rng}=A{r})*(Vagas!$B$2:$B${fim}="Líder"),0),0))&"","")'])
r = res.max_row + 1
res.append(["Total nas equipes fixas", f"=SUM(B3:B{r-1})", f"=SUM(C3:C{r-1})", f"=SUM(D3:D{r-1})", f"=SUM(E3:E{r-1})", ""])
for c in res[r]:
    c.font = Font(bold=True)
res["A" + str(r + 2)] = "Arrecadação e Captação inclui todos; a linha dela conta só a vaga de líder."
res["A" + str(r + 2)].alignment = Alignment(wrap_text=False)
wb.save(os.path.join(SAIDA, "planilha-equipes.xlsx"))
print(f"ok: {total} vagas em {len(EQUIPES)} equipes fixas + {TRANSVERSAL['nome']} -> {os.path.relpath(SAIDA, RAIZ)}")
