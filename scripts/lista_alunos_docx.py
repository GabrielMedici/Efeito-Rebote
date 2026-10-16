#!/usr/bin/env python3
"""Lista de alunos em Word (.docx), separada por turno (Noturno B e Matutino B), em ordem alfabética.

Modos:
  completo  Nº, Nome, CPF, RG e RA (relação para o presídio; dados pessoais, uso restrito)
  ra        Nome e RA (lista que a prof.ª Camila pediu em 08/10)

Uso (a saída deve ficar FORA do repositório, que é público):
  python3 scripts/lista_alunos_docx.py <dados.json> <saida.docx> [completo|ra]
"""
import json, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lista_alunos import fora_do_repo

TURMAS = [("Turma B, Noturno", "Noturno B"), ("Turma B, Matutino", "Matutino B"), ("Turma a confirmar", "")]


def main(dados, saida, modo="completo"):
    saida = fora_do_repo(saida)
    d = json.load(open(dados, encoding="utf-8"))
    al = sorted(d["alunos"], key=lambda a: a["nome"].casefold())
    completo = modo == "completo"
    doc = Document()
    s = doc.sections[0]
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10 if completo else 11)
    if completo:
        doc.add_heading("Efeito Rebote: relação de alunos para a visita", 1)
        doc.add_paragraph(f"Atualizado em {d.get('atualizado', '')}. Total: {len(al)} alunos. "
                          "Documento com dados pessoais (LGPD): uso restrito à organização da visita.")
        cab, larg = ["Nº", "Nome", "CPF", "RG", "RA"], [Cm(1), Cm(8.4), Cm(3.2), Cm(2.8), Cm(2.6)]
    else:
        doc.add_heading("Efeito Rebote: alunos por turma (nome e RA)", 1)
        cab, larg = ["Nome", "RA"], [Cm(12), Cm(5)]
    for tit, t in TURMAS:
        g = [a for a in al if (a.get("turma") or "") == t]
        if not g:
            continue
        doc.add_heading(f"{tit} ({len(g)} alunos)", 2)
        tb = doc.add_table(rows=1, cols=len(cab)); tb.style = "Table Grid"
        for i, h in enumerate(cab):
            c = tb.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
        for n, a in enumerate(g, 1):
            vals = [str(n), a["nome"], a["cpf"], a["rg"], a["ra"]] if completo else [a["nome"], a["ra"]]
            for i, v in enumerate(tb.add_row().cells):
                v.text = vals[i] or "[PENDENTE]"
        for row in tb.rows:
            for i, c in enumerate(row.cells):
                c.width = larg[i]
        print(tit, len(g))
    doc.save(saida)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(*sys.argv[1:4])
