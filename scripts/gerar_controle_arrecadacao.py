"""Gera entregas/acao-arrecadacao/controle-acao-arrecadacao.xlsx: repasse de R$ 150 por bloco, bilhetes (canhotos),
despesas e resumo. Pode ser importada no Google Planilhas. Uso: python3 scripts/gerar_controle_arrecadacao.py"""
import os
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.join(os.path.dirname(__file__), "..")
TOTAL, POR_ALUNO, ALUNOS, PRECO, PREMIO = 2400, 30, 80, 5, 1600
CAB = PatternFill("solid", fgColor="1B3A8C")
BRANCO = Font(bold=True, color="FFFFFF")


def cabecalho(ws, cols, larg):
    ws.append(cols)
    for i, (c, w) in enumerate(zip(ws[1], larg), 1):
        c.fill, c.font, c.alignment = CAB, BRANCO, Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


VALOR_BLOCO = POR_ALUNO * PRECO
VERDE, AMARELO = PatternFill("solid", fgColor="D9F2D9"), PatternFill("solid", fgColor="FFF3CD")
REAIS = '"R$" #,##0.00'

wb = Workbook()
wb.properties.creator = wb.properties.lastModifiedBy = "Acadêmicos do 3º semestre noturno – Turma B"
res = wb.active
res.title = "Resumo"
rep = wb.create_sheet("Repasses")
bil = wb.create_sheet("Bilhetes")
desp = wb.create_sheet("Despesas")

# Repasses: uma linha por acadêmico; o bloco fica quitado com R$ 150 identificados no extrato e canhotos conferidos
cabecalho(rep, ["Aluno nº", "Nome", "Período", "Bilhetes", "Data do repasse", "Valor repassado (R$)",
                "Canhotos entregues?", "Situação"], (9, 34, 12, 13, 15, 19, 18, 14))
n = ALUNOS + 1
for a in range(1, ALUNOS + 1):
    r = a + 1
    rep.append([a, "", "", f"{(a-1)*POR_ALUNO+1:04d}–{a*POR_ALUNO:04d}", "", None, "",
                f'=IF(AND(F{r}>={VALOR_BLOCO},G{r}="Sim"),"Quitado",IF(F{r}>0,"Parcial","Pendente"))'])
    rep[f"F{r}"].number_format = REAIS
    rep[f"E{r}"].number_format = "DD/MM/YYYY"
for col, op in (("C", '"Matutino,Noturno"'), ("G", '"Sim,Não"')):
    v = DataValidation(type="list", formula1=op, allow_blank=True)
    rep.add_data_validation(v)
    v.add(f"{col}2:{col}{n}")
rep.conditional_formatting.add(f"H2:H{n}", CellIsRule(operator="equal", formula=['"Quitado"'], fill=VERDE))
rep.conditional_formatting.add(f"H2:H{n}", CellIsRule(operator="equal", formula=['"Parcial"'], fill=AMARELO))

# Bilhetes: transcrição dos canhotos, que forma a lista do sorteio (dados pessoais só para contato com o ganhador)
cabecalho(bil, ["Bilhete", "Aluno nº", "Comprador", "Telefone", "Concorre?"], (9, 9, 30, 16, 11))
for b in range(1, TOTAL + 1):
    r = b + 1
    a = (b - 1) // POR_ALUNO + 1
    bil.append([f"{b:04d}", a, "", "", f'=IF(INDEX(Repasses!H:H,B{r}+1)="Quitado","Sim","Não")'])

cabecalho(desp, ["Data", "Descrição", "Fornecedor", "Nº da nota fiscal", "Valor (R$)"], (12, 36, 26, 16, 14))
for r in range(2, 202):
    desp[f"E{r}"].number_format = REAIS
    desp[f"A{r}"].number_format = "DD/MM/YYYY"

linhas = [
    ("Acadêmicos (blocos)", ALUNOS),
    ("Valor por bloco (R$)", VALOR_BLOCO),
    ("Blocos quitados", f'=COUNTIF(Repasses!H2:H{n},"Quitado")'),
    ("Blocos parciais", f'=COUNTIF(Repasses!H2:H{n},"Parcial")'),
    ("Blocos pendentes", f'=COUNTIF(Repasses!H2:H{n},"Pendente")'),
    ("Total repassado (R$)", f"=SUM(Repasses!F2:F{n})"),
    ("Meta (todos os blocos) (R$)", f"={ALUNOS}*{VALOR_BLOCO}"),
    ("% da meta", "=B7/B8"),
    ("Despesas com nota fiscal (R$)", "=SUM(Despesas!E2:E201)"),
    ("Saldo que deve constar no extrato (R$)", "=B7-B10"),
    ("Custo estimado do prêmio (R$)", PREMIO),
    ("Compra do prêmio liberada? (≥ 2× o custo)", '=IF(B7>=2*B12,"SIM","Ainda não")'),
    ("Bilhetes que concorrem", f'=COUNTIF(Bilhetes!E2:E{TOTAL+1},"Sim")'),
    ("Prazo de repasse e canhotos", "30/10/2026, 23h59"),
    ("Chave PIX (aleatória)", "fe5450d9-8b9e-470e-8b46-0d07e4a86d0d"),
    ("Sorteio", "02/11/2026, ao vivo @efeitorebote.oficial"),
]
res.append(["Ação de arrecadação solidária Efeito Rebote: resumo", ""])
res["A1"].font = Font(bold=True, size=13, color="1B3A8C")
for rot, val in linhas:
    res.append([rot, val])
res.column_dimensions["A"].width, res.column_dimensions["B"].width = 44, 40
res["B9"].number_format = "0.0%"
for c in ("B3", "B7", "B8", "B10", "B11", "B12"):
    res[c].number_format = REAIS

destino = os.path.join(RAIZ, "entregas", "acao-arrecadacao", "controle-acao-arrecadacao.xlsx")
wb.save(destino)
print("ok", destino)
