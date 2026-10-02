"""Gera entregas/rifa/controle-rifa.xlsx (distribuição, vendas, balanço por aluno e resumo).
Importe no Google Planilhas e ligue ao formulário de vendas. Uso: python3 scripts/gerar_controle_rifa.py"""
import os
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.join(os.path.dirname(__file__), "..")
TOTAL, POR_ALUNO, ALUNOS, PRECO, PREMIO = 2400, 30, 80, 5, 1600
CAB = PatternFill("solid", fgColor="1B3A8C")
BRANCO = Font(bold=True, color="FFFFFF")


def cabecalho(ws, cols, larg):
    ws.append(cols)
    for i, (c, w) in enumerate(zip(ws[1], larg), 1):
        c.fill, c.font, c.alignment = CAB, BRANCO, Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


wb = Workbook()
wb.properties.creator = wb.properties.lastModifiedBy = "Acadêmicos do 3º semestre noturno – Turma B"
res = wb.active
res.title = "Resumo"
dist = wb.create_sheet("Distribuição")
vend = wb.create_sheet("Vendas")
bal = wb.create_sheet("Por aluno")

cabecalho(dist, ["Aluno nº", "Nome", "Período", "Bilhete inicial", "Bilhete final", "Assinatura de recebimento"], (9, 34, 12, 14, 13, 26))
for a in range(1, ALUNOS + 1):
    dist.append([a, "", "", f"{(a-1)*POR_ALUNO+1:04d}", f"{a*POR_ALUNO:04d}", ""])
dv = DataValidation(type="list", formula1='"Matutino,Noturno"', allow_blank=True)
dist.add_data_validation(dv)
dv.add(f"C2:C{ALUNOS+1}")

cabecalho(vend, ["Bilhete", "Aluno nº", "Vendedor(a)", "Comprador", "Telefone", "Pagamento", "Pago?", "Data", "Conferido (comissão)"], (9, 9, 26, 28, 16, 12, 8, 12, 18))
for n in range(1, TOTAL + 1):
    r = n + 1
    vend.append([f"{n:04d}", (n - 1) // POR_ALUNO + 1, f"=IF(VLOOKUP(B{r},'Distribuição'!A:B,2,FALSE)=\"\",\"\",VLOOKUP(B{r},'Distribuição'!A:B,2,FALSE))", "", "", "", "", "", ""])
for col, op in (("F", '"PIX,Dinheiro"'), ("G", '"Sim,Não"'), ("I", '"Sim,Não"')):
    v = DataValidation(type="list", formula1=op, allow_blank=True)
    vend.add_data_validation(v)
    v.add(f"{col}2:{col}{TOTAL+1}")
vend.conditional_formatting.add(f"G2:G{TOTAL+1}", CellIsRule(operator="equal", formula=['"Sim"'], fill=PatternFill("solid", fgColor="D9F2D9")))

cabecalho(bal, ["Aluno nº", "Nome", "Vendidos", "Pagos", "Valor pago (R$)", "Restantes"], (9, 34, 10, 8, 15, 10))
for a in range(1, ALUNOS + 1):
    r = a + 1
    bal.append([a, f"=IF('Distribuição'!B{r}=\"\",\"\",'Distribuição'!B{r})", f'=COUNTIFS(Vendas!B:B,A{r},Vendas!D:D,"<>")',
                f'=COUNTIFS(Vendas!B:B,A{r},Vendas!G:G,"Sim")', f"=D{r}*{PRECO}", f"={POR_ALUNO}-C{r}"])

linhas = [
    ("Bilhetes no total", TOTAL),
    ("Bilhetes vendidos (com comprador)", '=COUNTIF(Vendas!D2:D2401,"<>")'),
    ("Bilhetes pagos", '=COUNTIF(Vendas!G2:G2401,"Sim")'),
    ("Arrecadado (R$)", f"=B4*{PRECO}"),
    ("% vendido (pago)", "=B4/B2"),
    ("Custo estimado do prêmio (R$)", PREMIO),
    ("Ponto de equilíbrio (bilhetes)", f"=B7/{PRECO}"),
    ("Compra do prêmio liberada? (≥ 2× equilíbrio)", '=IF(B4>=2*B8,"SIM","Ainda não")'),
    ("Saldo para itens após o prêmio (R$)", "=B5-B7"),
    ("PIX recebidos (R$)", f'=COUNTIFS(Vendas!F2:F2401,"PIX",Vendas!G2:G2401,"Sim")*{PRECO}'),
    ("Dinheiro recebido (R$)", f'=COUNTIFS(Vendas!F2:F2401,"Dinheiro",Vendas!G2:G2401,"Sim")*{PRECO}'),
    ("Prazo final de vendas", "30/10/2026 23h59"),
    ("Sorteio", "02/11/2026, ao vivo @efeitorebote.oficial"),
]
res.append(["Rifa solidária Efeito Rebote: resumo", ""])
res["A1"].font = Font(bold=True, size=13, color="1B3A8C")
for rot, val in linhas:
    res.append([rot, val])
res.column_dimensions["A"].width, res.column_dimensions["B"].width = 44, 36
res["B6"].number_format = "0.0%"
for c in ("B5", "B7", "B10", "B11", "B12"):
    res[c].number_format = '"R$" #,##0.00'

destino = os.path.join(RAIZ, "entregas", "rifa", "controle-rifa.xlsx")
wb.save(destino)
print("ok", destino)
