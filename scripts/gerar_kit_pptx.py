"""Monta os PowerPoints editáveis do Kit da Comunicação a partir das cenas extraídas
(scripts/extrair_cena_kit.mjs). Cada post vira um .pptx de 1080 × 1350 px com texto, formas e
imagens NATIVOS (o texto se edita; nada é foto do slide), as fontes Barlow e a legenda nas notas.
Importa no Canva (Criar um design > Importar arquivo), no PowerPoint e no Google Slides.
Uso: python3 scripts/gerar_kit_pptx.py <pasta das cenas> <pasta de saída> [slugs separados por vírgula]
1 px = 9525 EMU = 0,75 pt."""
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525
# Ajuste vertical do texto. O navegador centra a altura do texto (1,2 x o tamanho nas duas Barlow) dentro da
# entrelinha; o PowerPoint com entrelinha exata encosta o texto embaixo. A diferença é (1,2 x tamanho - entrelinha) / 2,
# mais cerca de 1 px medido (comparação pixel a pixel entre o LibreOffice e o PNG do navegador, em 07/10).
ALTURA_TEXTO = 1.2
FOLGA_PX = 1.0
# Modo de entrelinha (teste): "lo" = correção do LibreOffice; "canva" = sinal invertido; "pct" = entrelinha proporcional sem correção.
MODO = os.environ.get("ER_PPTX_MODO", "lo")


def px(v):
    return Emu(int(round(v * PX)))


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def fonte(familia, peso):
    """(nome da família no arquivo da fonte, negrito). Os pesos 500 e 600 são famílias próprias."""
    if "Condensed" in familia:
        return ("Barlow Condensed", True) if peso >= 700 else ("Barlow Condensed SemiBold", False)
    if peso >= 700:
        return "Barlow", True
    if peso >= 600:
        return "Barlow SemiBold", False
    if peso >= 500:
        return "Barlow Medium", False
    return "Barlow", False


def sem_estilo(shape):
    """Tira o <p:style> padrão (sombra e contorno do tema): cada forma leva só o que declaramos."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def quebrar(runs, linhas):
    """Divide os runs nas quebras de linha que o navegador fez (conta letras sem espaço)."""
    alvo = [len(re.sub(r"\s", "", l)) for l in linhas]
    saida, k, cont = [[]], 0, 0
    for run in runs:
        buf = ""
        for ch in run["t"]:
            if k < len(alvo) - 1 and cont >= alvo[k]:
                if not ch.strip():
                    continue  # espaço no ponto da quebra: some
                if buf:
                    saida[-1].append((run, buf)); buf = ""
                saida.append([]); k += 1; cont = 0
            buf += ch
            if ch.strip():
                cont += 1
        if buf:
            saida[-1].append((run, buf))
    for ln in saida:  # sem espaços nas pontas de cada linha
        if ln:
            ln[0] = (ln[0][0], ln[0][1].lstrip()); ln[-1] = (ln[-1][0], ln[-1][1].rstrip())
    return [ln for ln in saida if ln]


def caixa_texto(slide, it):
    runs = it["runs"]
    nlin = it["linhas"]
    alt = it["alturaLinha"]
    quebra_manual = nlin > 1 and it.get("equilibrar")
    larg_auto = None
    if nlin > 1 and not quebra_manual:
        # Largura que preserva as quebras do navegador em qualquer programa: cabe a linha mais larga e NÃO cabe a
        # próxima palavra. Fica no meio da faixa, para tolerar pequenas diferenças de medida da fonte.
        S = it["runs"][0]["tam"]
        piso = max(it["largLinhas"])
        teto = min(it["largLinhas"][i] + 0.26 * S + it["largPrimeiraPalavra"][i + 1] for i in range(nlin - 1))
        if teto - piso >= 8:
            larg_auto = (piso + teto) / 2
        else:
            quebra_manual = True  # faixa estreita demais: fixa as quebras
    nowrap = nlin == 1 or quebra_manual
    folga = 1.10 if nowrap else 1.0
    larg_uniao = it["w"]
    if nowrap:
        w = larg_uniao * folga + 8
        if it["alinhar"] == "center":
            x = it["x"] + larg_uniao / 2 - w / 2
        elif it["alinhar"] == "right":
            x = it["x"] + larg_uniao - w
        else:
            x = it["x"]
    else:
        w = larg_auto
        x = it["x"]
    fam0, _ = fonte(runs[0]["familia"], runs[0]["peso"])
    delta = (ALTURA_TEXTO * runs[0]["tam"] - alt) / 2
    y = it["y"] + {"lo": delta + FOLGA_PX, "canva": -delta + FOLGA_PX, "pct": 0}[MODO]
    h = nlin * alt
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tb.name = it["nome"]
    tf = tb.text_frame
    tf.word_wrap = not nowrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    par = tf.paragraphs[0]
    par.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(it["alinhar"], PP_ALIGN.LEFT)
    pPr = par._p.get_or_add_pPr()
    for tag in ("a:lnSpc", "a:spcBef", "a:spcAft"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    ln = pPr.makeelement(qn("a:lnSpc"), {})
    if MODO == "pct":
        sp = ln.makeelement(qn("a:spcPct"), {"val": str(int(round(alt / (ALTURA_TEXTO * runs[0]["tam"]) * 100000)))})
    else:
        sp = ln.makeelement(qn("a:spcPts"), {"val": str(int(round(alt * 0.75 * 100)))})
    ln.append(sp)
    pPr.insert(0, ln)
    for tag in ("a:spcBef", "a:spcAft"):
        e = pPr.makeelement(qn(tag), {}); e.append(e.makeelement(qn("a:spcPts"), {"val": "0"})); pPr.append(e)

    def escrever(run, texto):
        r = par.add_run(); r.text = texto
        nome, neg = fonte(run["familia"], run["peso"])
        f = r.font; f.name = nome; f.bold = neg; f.italic = bool(run.get("italico")); f.size = Pt(run["tam"] * 0.75); f.color.rgb = rgb(run["cor"])
        rPr = r._r.get_or_add_rPr(); rPr.set("lang", "pt-BR")
        if abs(run["espaco"]) > 0.01:
            rPr.set("spc", str(int(round(run["espaco"] * 0.75 * 100))))
        for tag in ("a:ea", "a:cs"):  # a mesma família nos três slots, para o Canva não trocar
            e = rPr.makeelement(qn(tag), {"typeface": nome}); rPr.append(e)

    if quebra_manual or (nlin > 1 and it.get("quebraManual")):
        for i, linha in enumerate(quebrar(runs, it["linhasTexto"])):
            if i:
                par.add_line_break()
            for run, texto in linha:
                escrever(run, texto)
    else:
        for run in runs:
            escrever(run, run["t"])
    return tb


def forma(slide, it):
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    b = it.get("borda")
    if b:  # a borda do CSS fica por dentro da caixa; a do PowerPoint é centrada na linha
        x += b["w"] / 2; y += b["w"] / 2; w -= b["w"]; h -= b["w"]
    if it["elipse"]:
        tipo = MSO_SHAPE.OVAL
    elif it["raio"] >= 0.5:
        tipo = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        tipo = MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(tipo, px(x), px(y), px(w), px(h))
    sem_estilo(sh)
    sh.name = it["nome"]
    if tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, max(0.0, (it["raio"] - (b["w"] / 2 if b else 0)) / min(w, h)))
    if it["fundo"]:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(it["fundo"])
    else:
        sh.fill.background()
    if b:
        sh.line.color.rgb = rgb(b["cor"]); sh.line.width = Pt(b["w"] * 0.75)
    else:
        sh.line.fill.background()
    return sh


def linha(slide, it):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(it["x"]), px(it["y"]), px(it["x"] + it["w"]), px(it["y"]))
    sem_estilo(c)
    c.name = it["nome"]; c.line.color.rgb = rgb(it["cor"]); c.line.width = Pt(it["espessura"] * 0.75)
    return c


def imagem(slide, it, pasta):
    pic = slide.shapes.add_picture(os.path.join(pasta, it["arquivo"]), px(it["x"]), px(it["y"]), px(it["w"]), px(it["h"]))
    pic.name = it["nome"]
    pic._element.nvPicPr.cNvPr.set("descr", {"Selo do projeto": "Selo do projeto Efeito Rebote", "Seta de rebote": "Seta de rebote, marca do projeto"}.get(it["nome"], "Seta"))
    return pic


def tema_fontes(prs):
    """Fonte do tema = Barlow, para qualquer texto novo já nascer na fonte da marca."""
    parte = prs.slide_master.part.part_related_by(RT.THEME)
    xml = parte.blob.decode("utf-8")
    xml = re.sub(r"(<a:majorFont>\s*<a:latin typeface=\")[^\"]*(\")", r"\1Barlow Condensed\2", xml)
    xml = re.sub(r"(<a:minorFont>\s*<a:latin typeface=\")[^\"]*(\")", r"\1Barlow\2", xml)
    xml = xml.replace('name="Office Theme"', 'name="Efeito Rebote"')
    parte._blob = xml.encode("utf-8")


def montar(cena_dir, saida):
    cena = json.load(open(os.path.join(cena_dir, "cena.json"), encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1080), px(1350)
    tema_fontes(prs)
    cp = prs.core_properties
    cp.title = f"Efeito Rebote — {cena['rot']}: {cena['titulo']}"; cp.author = "Efeito Rebote · Comunicação"; cp.subject = "Kit da Comunicação"; cp.language = "pt-BR"
    for i, sd in enumerate(cena["slides"]):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        if sd.get("fundoImagem"):
            bg = s.shapes.add_picture(os.path.join(cena_dir, sd["fundoImagem"]), 0, 0, px(1080), px(1350)); bg.name = "Fundo (foto de referência)"
            bg._element.nvPicPr.cNvPr.set("descr", "Fundo listrado: troque pela foto real")
        elif sd["bg"]:
            s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(sd["bg"])
        for it in sd["itens"]:
            if it["tipo"] == "forma":
                forma(s, it)
            elif it["tipo"] == "linha":
                linha(s, it)
            elif it["tipo"] == "imagem":
                imagem(s, it, cena_dir)
            elif it["tipo"] == "texto":
                caixa_texto(s, it)
        notas = s.notes_slide.notes_text_frame
        notas.text = (f"{cena['kicker']}\n{cena['titulo']}\n\nLEGENDA PRONTA\n{cena['legenda']}" if i == 0 else f"Slide {i + 1} de {len(cena['slides'])}")
    os.makedirs(saida, exist_ok=True)
    arq = os.path.join(saida, f"efeito-rebote-{cena['slug']}.pptx")
    prs.save(arq)
    return arq


if __name__ == "__main__":
    base, saida = sys.argv[1], sys.argv[2]
    so = sys.argv[3].split(",") if len(sys.argv) > 3 else None
    for nome in sorted(os.listdir(base)):
        if os.path.isfile(os.path.join(base, nome, "cena.json")) and (not so or nome in so):
            print("ok", montar(os.path.join(base, nome), saida))
